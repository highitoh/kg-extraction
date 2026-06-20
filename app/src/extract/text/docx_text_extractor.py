import os
from typing import Any, Dict

from docx import Document
from langchain_core.runnables.config import RunnableConfig

from .text_extractor import TextExtractor


class DocxTextExtractor(TextExtractor):
    """Extract text and tables from DOCX files"""

    def __init__(self):
        super().__init__("/workspace/app/log/docx_text_extractor")

    def _extract_text(self, docx_path: str) -> str:
        """
        Extract text from DOCX using python-docx

        Returns:
            Concatenated paragraph text
        """
        if not os.path.exists(docx_path):
            raise FileNotFoundError(f"DOCX file not found: {docx_path}")

        if not docx_path.lower().endswith(".docx"):
            raise ValueError("Not a DOCX file")

        doc = Document(docx_path)
        text_parts = []

        for paragraph in doc.paragraphs:
            para_text = paragraph.text.strip()
            if para_text:  # Skip empty paragraphs
                text_parts.append(para_text)

        return "\n".join(text_parts)

    def _extract_table_blocks(self, docx_path: str) -> list:
        """
        Extract structured table blocks from DOCX

        Returns:
            List of structured blocks (same format as PDF tables)
        """
        from src.utils.structured_block_builder import build_table_blocks
        from src.utils.table_extractor import extract_tables_from_docx

        all_blocks = []
        doc = Document(docx_path)

        for table_idx, table in enumerate(doc.tables, start=1):
            # Extract and normalize table
            tables = extract_tables_from_docx(table)
            if not tables:
                continue

            # Build metadata
            filename = os.path.basename(docx_path)
            file_id = f"docx:{filename}:table_{table_idx}"
            scope_ref = f"table_{table_idx}"
            section_path = f"{filename} - Table {table_idx}"

            # Extract context (paragraphs before/after table)
            context_window = self._get_table_context(doc, table, max_chars=200)

            # Build structured blocks
            table_blocks = build_table_blocks(tables, file_id, scope_ref, section_path, context_window)
            all_blocks.extend(table_blocks)

        return all_blocks

    def _get_table_context(self, doc: Document, table, max_chars: int) -> str:
        """
        Extract surrounding text context for a table

        Looks for paragraphs before and after the table to provide context
        """
        # Find table position in document
        table_elements = [elem for elem in doc.element.body]
        try:
            table_idx = table_elements.index(table._element)
        except ValueError:
            return ""

        # Get previous paragraph
        context_parts = []
        for i in range(table_idx - 1, -1, -1):
            elem = table_elements[i]
            if elem.tag.endswith("p"):  # Paragraph
                text = elem.text.strip() if hasattr(elem, "text") else ""
                if text:
                    context_parts.insert(0, text)
                    break

        # Get next paragraph if we need more context
        if sum(len(p) for p in context_parts) < max_chars:
            for i in range(table_idx + 1, len(table_elements)):
                elem = table_elements[i]
                if elem.tag.endswith("p"):
                    text = elem.text.strip() if hasattr(elem, "text") else ""
                    if text:
                        context_parts.append(text)
                        break

        context = " ".join(context_parts)
        return context[:max_chars] if context else ""

    def invoke(self, input: Dict[str, Any], config: RunnableConfig = None) -> Dict[str, Any]:
        """
        Main extraction entry point

        Input: {"files": [{"path": "/path/to/file.docx"}]}
        Output: {
            "file_id": "docx:filename.docx",
            "file_name": "filename.docx",
            "sentences": [{"text": "paragraph 1"}, ...],
            "_docx_structured_blocks": [...]  # Optional
        }
        """
        file_path = input["files"][0]["path"]

        # Extract text
        extracted_text = self._extract_text(file_path)

        # Extract structured table blocks
        all_blocks = self._extract_table_blocks(file_path)

        # Split text into sentences (one paragraph = one sentence)
        sentences = []
        for line_text in extracted_text.split("\n"):
            if line_text.strip():
                sentences.append({"text": line_text.strip()})

        # Build output using base class helper
        output = self.build_output(file_path, sentences)

        # Attach structured blocks if any tables exist
        if all_blocks:
            output["_docx_structured_blocks"] = all_blocks

        # Save log
        self.logger.save_log(output, "docx_text_extractor_output_")

        return output


if __name__ == "__main__":
    extractor = DocxTextExtractor()
    result = extractor.invoke({"files": [{"path": "../doc/sample.docx"}]})
    print(f"Processing result: {result}")
