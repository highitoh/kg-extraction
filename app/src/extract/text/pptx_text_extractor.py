import os
from typing import Any, Dict

from langchain_core.runnables.config import RunnableConfig
from pptx import Presentation

from .text_extractor import TextExtractor


class PptxTextExtractor(TextExtractor):
    """Extract text and tables from PPTX files"""

    def __init__(self):
        super().__init__("/workspace/app/log/pptx_text_extractor")

    def _extract_text(self, pptx_path: str) -> list[tuple[int, str]]:
        """
        Extract text from PPTX slide by slide using python-pptx

        Returns:
            list[tuple[int, str]]: List of (slide number, text) tuples
        """
        if not os.path.exists(pptx_path):
            raise FileNotFoundError(f"PPTX file not found: {pptx_path}")

        if not pptx_path.lower().endswith(".pptx"):
            raise ValueError("Not a PPTX file")

        slides = []
        try:
            prs = Presentation(pptx_path)
            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_text_parts = []

                # Extract text from all shapes on the slide
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text_parts.append(shape.text.strip())

                # Combine all text from this slide
                if slide_text_parts:
                    slide_text = "\n".join(slide_text_parts)
                    slides.append((slide_num, slide_text))

        except Exception as e:
            print(f"Error occurred during text extraction from PPTX: {e}")

        return slides

    def _extract_table_blocks(self, pptx_path: str) -> list:
        """
        Extract structured table blocks from PPTX

        Returns:
            List of structured blocks (same format as PDF/DOCX tables)
        """
        from src.utils.structured_block_builder import build_table_blocks
        from src.utils.table_extractor import extract_tables_from_pptx

        all_blocks = []
        prs = Presentation(pptx_path)
        filename = os.path.basename(pptx_path)

        for slide_num, slide in enumerate(prs.slides, start=1):
            table_idx_on_slide = 0

            for shape in slide.shapes:
                if shape.has_table:
                    table_idx_on_slide += 1

                    # Extract and normalize table
                    tables = extract_tables_from_pptx(shape.table)
                    if not tables:
                        continue

                    # Build metadata
                    file_id = f"pptx:{filename}:slide_{slide_num}:table_{table_idx_on_slide}"
                    scope_ref = f"slide_{slide_num}"
                    section_path = f"{filename} - Slide {slide_num}"

                    # Extract context (other text on the same slide)
                    context_window = self._get_slide_context(slide, max_chars=200)

                    # Build structured blocks
                    table_blocks = build_table_blocks(tables, file_id, scope_ref, section_path, context_window)
                    all_blocks.extend(table_blocks)

        return all_blocks

    def _get_slide_context(self, slide, max_chars: int) -> str:
        """
        Extract text context from the same slide

        Gets all non-table text from the slide to provide context
        """
        context_parts = []

        for shape in slide.shapes:
            # Skip table shapes
            if shape.has_table:
                continue

            # Extract text from shape
            if hasattr(shape, "text") and shape.text.strip():
                context_parts.append(shape.text.strip())

        context = " ".join(context_parts)
        return context[:max_chars] if context else ""

    def invoke(self, input: Dict[str, Any], config: RunnableConfig = None) -> Dict[str, Any]:
        """
        Main extraction entry point

        Input: {"files": [{"path": "/path/to/file.pptx"}]}
        Output: {
            "file_id": "pptx:filename.pptx:hash",
            "file_name": "filename.pptx",
            "sentences": [{"text": "[[Slide1]] ..."}, ...],
            "_pptx_structured_blocks": [...]  # Optional
        }
        """
        file_path = input["files"][0]["path"]

        # Extract text slide by slide
        slides = self._extract_text(file_path)

        # Extract structured table blocks
        all_blocks = self._extract_table_blocks(file_path)

        # Generate sentences with [[SlideN]] prefix
        sentences = []
        for slide_num, slide_text in slides:
            for line_text in slide_text.split("\n"):
                if line_text.strip():
                    sentences.append({"text": f"[[Slide{slide_num}]] {line_text.strip()}"})

        # Build output using base class helper
        output = self.build_output(file_path, sentences)

        # Attach structured blocks if any tables exist
        if all_blocks:
            output["_pptx_structured_blocks"] = all_blocks

        # Save log
        self.logger.save_log(output, "pptx_text_extractor_output_")

        return output


if __name__ == "__main__":
    extractor = PptxTextExtractor()
    result = extractor.invoke({"files": [{"path": "../doc/sample.pptx"}]})
    print(f"Processing result: {result}")
