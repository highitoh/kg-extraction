"""Image sources for ImageAnalyzer.

Each function turns one origin (PDF XObjects, rendered PDF pages, Office
documents, ...) into ImageSubject records. Anything origin-specific -- rule-based
filtering, page rasterisation, provenance fields -- belongs here rather than in
the analyzer.
"""

import io
import os
import uuid
from typing import Any, Dict, List, Optional, Set

import pdfplumber
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.extract.image.image_analyzer import ImageSubject
from src.utils.image_utils import extract_image_bytes, passes_basic_filters
from src.utils.pdf_page_to_image import convert_pdf_page_to_image


def embedded_images(file_path: str, file_id: str, progress: bool = True) -> List[ImageSubject]:
    """Collect embedded images from a file, picking the reader by extension.

    Formats without embedded images (TXT, Markdown) and URLs yield no subjects
    rather than an error, so callers can run this on any source.
    """
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return pdf_embedded_images(file_path, file_id, progress=progress)
    if extension == ".docx":
        return docx_embedded_images(file_path, file_id, progress=progress)
    if extension == ".pptx":
        return pptx_embedded_images(file_path, file_id, progress=progress)

    if progress:
        print(f"[embedded_images] No image reader for {extension or file_path!r}, skipping")
    return []


def pdf_embedded_images(file_path: str, file_id: str, progress: bool = True) -> List[ImageSubject]:
    """Collect images embedded in a PDF, skipping tiny and decorative ones.

    Args:
        file_path: Path to the PDF
        file_id: File ID carried into every chunk
        progress: Print a summary of what was filtered out

    Returns:
        One ImageSubject per image that passed the rule-based filters
    """
    subjects: List[ImageSubject] = []
    total_seen = 0
    rule_filtered = 0
    byte_extraction_failed = 0

    if progress:
        print(f"[pdf_embedded_images] Scanning {file_path}")

    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)
        for page_number, page in enumerate(pdf.pages, start=1):
            for image_dict in page.images or []:
                total_seen += 1
                if not passes_basic_filters(image_dict):
                    rule_filtered += 1
                    continue
                image_bytes = extract_image_bytes(page, image_dict)
                if not image_bytes:
                    byte_extraction_failed += 1
                    continue

                image_id = f"image-{uuid.uuid4()}"
                subjects.append(
                    ImageSubject(
                        load=_bytes_loader(image_bytes),
                        file_id=file_id,
                        page_number=page_number,
                        source={"kind": "image", "page_number": page_number, "image_id": image_id},
                    )
                )

    if progress:
        print(
            f"[pdf_embedded_images] pages={page_count} total_images={total_seen} "
            f"rule_filtered={rule_filtered} byte_extraction_failed={byte_extraction_failed} "
            f"candidates={len(subjects)}"
        )
    return subjects


def pdf_page_images(
    pdf_path: str,
    item_page_mappings: List[Dict[str, Any]],
    file_id: str,
    max_pixels: int = 2073600,
) -> List[ImageSubject]:
    """Rasterise the PDF pages that were matched to Miro Frames/Diagrams.

    Args:
        pdf_path: Path to the PDF exported from Miro
        item_page_mappings: Page-to-item matches, each {"page_number", "item_id", "item_type"}
        file_id: File ID carried into every chunk
        max_pixels: Pixel cap per rendered page

    Returns:
        One ImageSubject per mapped page
    """
    subjects: List[ImageSubject] = []

    for mapping in item_page_mappings:
        page_number = mapping["page_number"]
        item_id = mapping["item_id"]
        subjects.append(
            ImageSubject(
                load=_page_loader(pdf_path, page_number, max_pixels),
                file_id=f"{file_id}:{item_id}",
                page_number=page_number,
                source={
                    "kind": "image",
                    "page_number": page_number,
                    "item_id": item_id,
                    "item_type": mapping["item_type"],
                },
            )
        )

    return subjects


def docx_embedded_images(file_path: str, file_id: str, progress: bool = True) -> List[ImageSubject]:
    """Collect images embedded in a DOCX, skipping tiny and decorative ones.

    DOCX has no fixed pagination, so subjects carry no page number.

    Args:
        file_path: Path to the DOCX
        file_id: File ID carried into every chunk
        progress: Print a summary of what was filtered out

    Returns:
        One ImageSubject per image that passed the rule-based filters
    """
    from docx import Document

    subjects: List[ImageSubject] = []
    total_seen = 0
    rule_filtered = 0

    if progress:
        print(f"[docx_embedded_images] Scanning {file_path}")

    for part in Document(file_path).part.package.image_parts:
        total_seen += 1
        image = part.image
        if not passes_basic_filters({"width": image.px_width, "height": image.px_height}):
            rule_filtered += 1
            continue

        subjects.append(_office_subject(part.blob, file_id, page_number=None))

    if progress:
        print(
            f"[docx_embedded_images] total_images={total_seen} rule_filtered={rule_filtered} candidates={len(subjects)}"
        )
    return subjects


def pptx_embedded_images(file_path: str, file_id: str, progress: bool = True) -> List[ImageSubject]:
    """Collect pictures placed on PPTX slides, skipping tiny and decorative ones.

    Walks the shapes on each slide (descending into groups) so every subject
    keeps the slide it appeared on; media referenced only by layouts or masters
    is left out.

    Args:
        file_path: Path to the PPTX
        file_id: File ID carried into every chunk
        progress: Print a summary of what was filtered out

    Returns:
        One ImageSubject per picture that passed the rule-based filters
    """
    from pptx import Presentation

    subjects: List[ImageSubject] = []
    total_seen = 0
    rule_filtered = 0
    unreadable = 0

    if progress:
        print(f"[pptx_embedded_images] Scanning {file_path}")

    presentation = Presentation(file_path)
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for picture in _iter_pictures(slide.shapes):
            total_seen += 1
            try:
                image = picture.image
                width, height = image.size
                image_bytes = image.blob
            except Exception:
                unreadable += 1
                continue

            if not passes_basic_filters({"width": width, "height": height}):
                rule_filtered += 1
                continue

            subjects.append(_office_subject(image_bytes, file_id, page_number=slide_number))

    if progress:
        print(
            f"[pptx_embedded_images] slides={len(presentation.slides)} total_images={total_seen} "
            f"rule_filtered={rule_filtered} unreadable={unreadable} candidates={len(subjects)}"
        )
    return subjects


def miro_rendered_regions(
    items: List[Dict[str, Any]],
    connectors: List[Dict[str, Any]],
    file_id: str,
    changed_item_ids: Optional[Set[str]] = None,
    progress: bool = True,
) -> List[ImageSubject]:
    """Draw a Miro board from its API data, one image per frame, diagram or cluster.

    Args:
        items: Board items, with diagram children already resolved
        connectors: Board connectors, drawn as arrows between items
        file_id: File ID prefix; each region appends its own id
        changed_item_ids: In incremental mode, keep only regions holding one of these
        progress: Print what was rendered

    Returns:
        One ImageSubject per rendered region
    """
    from src.extract.miro.miro_board_renderer import MiroBoardRenderer

    if progress:
        print(f"[miro_rendered_regions] Rendering {len(items)} items and {len(connectors)} connectors")

    regions = MiroBoardRenderer().render(items, connectors)

    subjects: List[ImageSubject] = []
    skipped = 0
    for region in regions:
        if changed_item_ids is not None and not (set(region.item_ids) & changed_item_ids):
            skipped += 1
            continue

        source: Dict[str, Any] = {
            "kind": "image",
            "item_id": region.region_id,
            "item_type": region.region_type,
        }
        if region.title:
            source["title"] = region.title

        subjects.append(
            ImageSubject(
                load=_region_loader(region),
                file_id=f"{file_id}:{region.region_id}",
                source=source,
            )
        )

    if progress:
        print(f"[miro_rendered_regions] regions={len(regions)} unchanged_skipped={skipped} candidates={len(subjects)}")
    return subjects


def _region_loader(region):
    def load() -> Image.Image:
        return region.image

    return load


def _iter_pictures(shapes):
    """Yield picture shapes, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape


def _office_subject(image_bytes: bytes, file_id: str, page_number: Optional[int]) -> ImageSubject:
    source: Dict[str, Any] = {"kind": "image", "image_id": f"image-{uuid.uuid4()}"}
    if page_number is not None:
        source["page_number"] = page_number
    return ImageSubject(
        load=_bytes_loader(image_bytes),
        file_id=file_id,
        page_number=page_number,
        source=source,
    )


def _bytes_loader(image_bytes: bytes):
    def load() -> Image.Image:
        return Image.open(io.BytesIO(image_bytes))

    return load


def _page_loader(pdf_path: str, page_number: int, max_pixels: int):
    def load() -> Image.Image:
        return convert_pdf_page_to_image(pdf_path, page_number, max_pixels=max_pixels)

    return load
